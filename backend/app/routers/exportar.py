# ============================================================
# Router: exportar.py
# Descripción: Endpoints para exportar datos a Excel
# Acceso: según rol del usuario
# ============================================================

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from app.database import get_connection
from app.auth import verificar_token
import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
import io
from datetime import datetime
from app.routers.admin_pais import (
    _calcular_resumen_dinamica, _verificar_dinamica_finalizada,
    _calcular_oportunidades_sin_explotar, _calcular_sobre_stock_departamento,
    _calcular_cobertura_tienda, _calcular_redistribucion_departamento,
    _calcular_forecast_demanda, _calcular_impacto_campana_cobertura
)


router = APIRouter(prefix="/exportar", tags=["Exportar"])


def obtener_moneda_pais(cursor, pais: str) -> str:
    """
    Devuelve el símbolo de moneda del país — cada país ya carga sus
    datos en su propia moneda (colón, quetzal, dólar), así que los
    reportes exportados deben mostrar el símbolo correcto, no un $
    fijo. Si el país no tiene moneda configurada o no se especificó
    país (reporte multi-país de superadmin), usa $ como respaldo.
    """
    if not pais:
        return "$"
    cursor.execute("SELECT moneda_simbolo FROM cat_paises WHERE codigo = ?", (pais,))
    fila = cursor.fetchone()
    return fila[0] if fila and fila[0] else "$"


# ── Estilos de Excel ──────────────────────────────────────
def estilo_header(ws, fila, columnas, color_fondo="1E2761"):
    """Aplica estilo profesional a la fila de encabezados"""
    fill = PatternFill(start_color=color_fondo, end_color=color_fondo, fill_type="solid")
    font = Font(color="FFFFFF", bold=True, size=11, name="Calibri")
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        bottom=Side(style="medium", color="FFFFFF")
    )
    for col, valor in enumerate(columnas, 1):
        cell = ws.cell(row=fila, column=col, value=valor)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
        cell.border = border

def estilo_fila(ws, fila, valores, par=True):
    """Aplica estilo a filas de datos con colores alternados"""
    color = "F4F7FB" if par else "FFFFFF"
    fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    font = Font(name="Calibri", size=10)
    alignment = Alignment(vertical="center", wrap_text=False)
    for col, valor in enumerate(valores, 1):
        cell = ws.cell(row=fila, column=col, value=valor)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment

def ajustar_columnas(ws):
    """Ajusta el ancho de columnas automáticamente"""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

def agregar_titulo(ws, titulo, subtitulo, color="1E2761"):
    """Agrega título y subtítulo al inicio del reporte"""
    ws.merge_cells('A1:J1')
    cell_titulo = ws['A1']
    cell_titulo.value = titulo
    cell_titulo.font = Font(name="Calibri", bold=True, size=14, color=color)
    cell_titulo.alignment = Alignment(horizontal="center", vertical="center")
    cell_titulo.fill = PatternFill(start_color="F4F7FB", end_color="F4F7FB", fill_type="solid")
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:J2')
    cell_sub = ws['A2']
    cell_sub.value = subtitulo
    cell_sub.font = Font(name="Calibri", size=10, color="8892A4", italic=True)
    cell_sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

# ── Endpoint: Exportar artículos en dinámica ─────────────
@router.get("/articulos")
def exportar_articulos(usuario=Depends(verificar_token)):
    """
    Exporta todos los artículos en dinámica a Excel.
    Filtra por país según el rol del usuario.
    categoria y precio_regular salen de articulos_maestro vía LEFT JOIN
    por SKU — ya no viven en dinamica_detalle.
    """
    conn = get_connection()
    cursor = conn.cursor()
    pais = usuario["pais"]
    rol = usuario["rol"]

    join_catalogo = """
        LEFT JOIN (
            SELECT sku, pais, MIN(categoria) as categoria, AVG(precio_unitario) as precio_unitario
            FROM articulos_maestro
            GROUP BY sku, pais
        ) am ON am.sku = dd.sku AND am.pais = dd.pais
    """

    if pais and rol != "superadmin":
        cursor.execute(f"""
            SELECT dd.sku, dd.descripcion, am.categoria,
                   am.precio_unitario, dd.precio_dinamica,
                   dd.ubicacion_dinamica, dd.ubicacion_tienda, dd.tienda, dd.tipo_campana, dd.pais,
                   d.nombre as dinamica, d.fecha_inicio, d.fecha_fin
            FROM dinamica_detalle dd
            INNER JOIN dinamicas d ON dd.dinamica_id = d.id
            {join_catalogo}
            WHERE dd.activo = 1 AND d.activa = 1 AND dd.pais = ?
            ORDER BY dd.tienda, dd.ubicacion_dinamica, dd.sku
        """, (pais,))
    else:
        cursor.execute(f"""
            SELECT dd.sku, dd.descripcion, am.categoria,
                   am.precio_unitario, dd.precio_dinamica,
                   dd.ubicacion_dinamica, dd.ubicacion_tienda, dd.tienda, dd.tipo_campana, dd.pais,
                   d.nombre as dinamica, d.fecha_inicio, d.fecha_fin
            FROM dinamica_detalle dd
            INNER JOIN dinamicas d ON dd.dinamica_id = d.id
            {join_catalogo}
            WHERE dd.activo = 1 AND d.activa = 1
            ORDER BY dd.pais, dd.tienda, dd.ubicacion_dinamica, dd.sku
        """)

    filas = cursor.fetchall()
    moneda = obtener_moneda_pais(cursor, pais)
    conn.close()

    # ── Crear Excel ───────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Artículos en Dinámica"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    pais_label = {"CR": "Costa Rica", "GT": "Guatemala", "SV": "El Salvador"}.get(pais, "Todos los países")

    agregar_titulo(
        ws,
        f"PromoIQ — Artículos en Dinámica",
        f"Generado el {fecha_reporte} | País: {pais_label} | Usuario: {usuario['nombre']}"
    )

    columnas = [
        "Código", "Descripción", "Categoría",
        "Precio Regular", "Precio Dinámica", "Ahorro",
        "Ubicación Dinámica", "Ubicación Tienda", "Tienda", "Tipo Campaña", "País",
        "Dinámica", "Fecha Inicio", "Fecha Fin"
    ]
    estilo_header(ws, 3, columnas)
    ws.row_dimensions[3].height = 22

    for i, fila in enumerate(filas, 4):
        precio_regular = float(fila[3]) if fila[3] else 0
        precio_dinamica = float(fila[4]) if fila[4] else 0
        valores = [
            str(fila[0]),
            str(fila[1]),
            str(fila[2]) if fila[2] else "-",
            precio_regular,
            precio_dinamica,
            round(precio_regular - precio_dinamica, 2),
            str(fila[5]) if fila[5] else "-",
            str(fila[6]) if fila[6] else "-",
            str(fila[7]) if fila[7] else "-",
            str(fila[8]) if fila[8] else "-",
            str(fila[9]) if fila[9] else "-",
            str(fila[10]),
            str(fila[11]),
            str(fila[12]),
        ]
        estilo_fila(ws, i, valores, par=(i % 2 == 0))

    # Formato de moneda en columnas de precio — símbolo según el
    # país del reporte, no fijo en dólares
    for row in ws.iter_rows(min_row=4, max_row=len(filas) + 3, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = f'"{moneda}"#,##0.00'

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    # ── Guardar y retornar ────────────────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"PromoIQ_Articulos_{pais_label}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )

# ── Endpoint: Exportar historial de consultas ─────────────
@router.get("/historial")
def exportar_historial(usuario=Depends(verificar_token)):
    """
    Exporta el historial de consultas a la IA en Excel.
    Solo disponible para admin y superadmin.
    """
    if usuario["rol"] not in ["superadmin", "admin"]:
        raise HTTPException(status_code=403, detail="Acceso no autorizado")

    conn = get_connection()
    cursor = conn.cursor()
    pais = usuario["pais"]
    rol = usuario["rol"]

    if pais and rol != "superadmin":
        cursor.execute("""
            SELECT u.nombre, r.nombre as rol, c.pregunta,
                   c.tokens_usados, c.fecha
            FROM consultas_log c
            INNER JOIN usuarios u ON c.usuario_id = u.id
            INNER JOIN cat_roles r ON u.rol_id = r.id
            WHERE u.pais = ?
            ORDER BY c.fecha DESC
        """, (pais,))
    else:
        cursor.execute("""
            SELECT u.nombre, r.nombre as rol, c.pregunta,
                   c.tokens_usados, c.fecha
            FROM consultas_log c
            INNER JOIN usuarios u ON c.usuario_id = u.id
            INNER JOIN cat_roles r ON u.rol_id = r.id
            ORDER BY c.fecha DESC
        """)

    filas = cursor.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial de Consultas"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    pais_label = {"CR": "Costa Rica", "GT": "Guatemala", "SV": "El Salvador"}.get(pais, "Todos los países")

    agregar_titulo(
        ws,
        "PromoIQ — Historial de Consultas IA",
        f"Generado el {fecha_reporte} | País: {pais_label} | Total: {len(filas)} consultas"
    )

    columnas = ["Usuario", "Rol", "Pregunta", "Tokens Usados", "Fecha"]
    estilo_header(ws, 3, columnas, color_fondo="0077B6")
    ws.row_dimensions[3].height = 22

    for i, fila in enumerate(filas, 4):
        valores = [
            str(fila[0]),
            str(fila[1]),
            str(fila[2]),
            int(fila[3]) if fila[3] else 0,
            str(fila[4]),
        ]
        estilo_fila(ws, i, valores, par=(i % 2 == 0))

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"PromoIQ_Historial_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )

# ── Endpoint: Exportar reporte completo ───────────────────
@router.get("/reporte-completo")
def exportar_reporte_completo(usuario=Depends(verificar_token)):
    """
    Exporta un reporte completo con múltiples hojas:
    - Artículos por dinámica
    - Resumen por ubicación
    - Resumen por tipo de campaña
    Solo disponible para admin y superadmin.
    categoria y precio_regular salen de articulos_maestro vía LEFT JOIN por SKU.
    """
    if usuario["rol"] not in ["superadmin", "admin"]:
        raise HTTPException(status_code=403, detail="Acceso no autorizado")

    conn = get_connection()
    cursor = conn.cursor()
    pais = usuario["pais"]
    rol = usuario["rol"]
    filtro = f"AND dd.pais = '{pais}'" if pais and rol != "superadmin" else ""
    pais_label = {"CR": "Costa Rica", "GT": "Guatemala", "SV": "El Salvador"}.get(pais, "Todos los países")
    moneda = obtener_moneda_pais(cursor, pais)
    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")

    join_catalogo = """
        LEFT JOIN (
            SELECT sku, pais, MIN(categoria) as categoria, AVG(precio_unitario) as precio_unitario
            FROM articulos_maestro
            GROUP BY sku, pais
        ) am ON am.sku = dd.sku AND am.pais = dd.pais
    """

    wb = openpyxl.Workbook()

    # ── Hoja 1: Artículos ─────────────────────────────────
    ws1 = wb.active
    ws1.title = "Artículos"
    agregar_titulo(ws1, "Artículos en Dinámica", f"Generado el {fecha_reporte} | {pais_label}")

    cursor.execute(f"""
        SELECT dd.sku, dd.descripcion, am.categoria,
               am.precio_unitario, dd.precio_dinamica,
               dd.ubicacion_dinamica, dd.tienda, dd.tipo_campana, d.nombre
        FROM dinamica_detalle dd
        INNER JOIN dinamicas d ON dd.dinamica_id = d.id
        {join_catalogo}
        WHERE dd.activo = 1 AND d.activa = 1 {filtro}
        ORDER BY dd.tienda, dd.ubicacion_dinamica
    """)
    articulos = cursor.fetchall()

    estilo_header(ws1, 3, ["Código", "Descripción", "Categoría", "P. Regular", "P. Dinámica", "Ahorro", "Ubicación", "Tienda", "Tipo Campaña", "Dinámica"])
    for i, f in enumerate(articulos, 4):
        precio_regular = float(f[3] or 0)
        precio_dinamica = float(f[4] or 0)
        estilo_fila(ws1, i, [
            str(f[0]), str(f[1]), str(f[2] or "-"),
            precio_regular, precio_dinamica, round(precio_regular - precio_dinamica, 2),
            str(f[5] or "-"), str(f[6] or "-"), str(f[7] or "-"), str(f[8])
        ], par=(i % 2 == 0))
    for row in ws1.iter_rows(min_row=4, max_row=len(articulos)+3, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = f'"{moneda}"#,##0.00'
    ajustar_columnas(ws1)
    ws1.freeze_panes = "A4"

    # ── Hoja 2: Por Ubicación ─────────────────────────────
    # El ahorro total ya no se puede sumar directo en SQL (precio_regular
    # vive en articulos_maestro, no en dinamica_detalle) — se agrega por
    # fila en Python después de traer el detalle con el JOIN.
    ws2 = wb.create_sheet("Por Ubicación")
    agregar_titulo(ws2, "Artículos por Ubicación", f"Generado el {fecha_reporte} | {pais_label}")

    cursor.execute(f"""
        SELECT dd.ubicacion_dinamica, dd.precio_dinamica, am.precio_unitario
        FROM dinamica_detalle dd
        INNER JOIN dinamicas d ON dd.dinamica_id = d.id
        {join_catalogo}
        WHERE dd.activo = 1 AND d.activa = 1 {filtro}
    """)
    filas_ubicacion = cursor.fetchall()

    resumen_ubicacion = {}
    for ubicacion, precio_dinamica, precio_regular in filas_ubicacion:
        clave = ubicacion or "-"
        if clave not in resumen_ubicacion:
            resumen_ubicacion[clave] = {"total": 0, "suma_precio": 0, "suma_ahorro": 0}
        precio_dinamica = float(precio_dinamica or 0)
        precio_regular = float(precio_regular or 0)
        resumen_ubicacion[clave]["total"] += 1
        resumen_ubicacion[clave]["suma_precio"] += precio_dinamica
        resumen_ubicacion[clave]["suma_ahorro"] += (precio_regular - precio_dinamica)

    ubicaciones = sorted(resumen_ubicacion.items(), key=lambda x: x[1]["total"], reverse=True)

    estilo_header(ws2, 3, ["Ubicación", "Total Artículos", "Precio Promedio", "Ahorro Total"], color_fondo="06847A")
    for i, (nombre_ubic, datos) in enumerate(ubicaciones, 4):
        precio_prom = datos["suma_precio"] / datos["total"] if datos["total"] > 0 else 0
        estilo_fila(ws2, i, [
            nombre_ubic, datos["total"], round(precio_prom, 2), round(datos["suma_ahorro"], 2)
        ], par=(i % 2 == 0))
    for row in ws2.iter_rows(min_row=4, max_row=len(ubicaciones)+3, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = f'"{moneda}"#,##0.00'
    ajustar_columnas(ws2)
    ws2.freeze_panes = "A4"

    # ── Hoja 3: Por Tipo de Campaña ───────────────────────
    ws3 = wb.create_sheet("Por Tipo Campaña")
    agregar_titulo(ws3, "Artículos por Tipo de Campaña", f"Generado el {fecha_reporte} | {pais_label}")

    cursor.execute(f"""
        SELECT dd.tipo_campana, dd.precio_dinamica, am.precio_unitario
        FROM dinamica_detalle dd
        INNER JOIN dinamicas d ON dd.dinamica_id = d.id
        {join_catalogo}
        WHERE dd.activo = 1 AND d.activa = 1 {filtro}
    """)
    filas_tipo = cursor.fetchall()

    resumen_tipo = {}
    for tipo_campana, precio_dinamica, precio_regular in filas_tipo:
        clave = tipo_campana or "-"
        if clave not in resumen_tipo:
            resumen_tipo[clave] = {"total": 0, "suma_precio": 0, "suma_ahorro": 0}
        precio_dinamica = float(precio_dinamica or 0)
        precio_regular = float(precio_regular or 0)
        resumen_tipo[clave]["total"] += 1
        resumen_tipo[clave]["suma_precio"] += precio_dinamica
        resumen_tipo[clave]["suma_ahorro"] += (precio_regular - precio_dinamica)

    tipos = sorted(resumen_tipo.items(), key=lambda x: x[1]["total"], reverse=True)

    estilo_header(ws3, 3, ["Tipo Campaña", "Total Artículos", "Precio Promedio", "Ahorro Total"], color_fondo="B45309")
    for i, (nombre_tipo, datos) in enumerate(tipos, 4):
        precio_prom = datos["suma_precio"] / datos["total"] if datos["total"] > 0 else 0
        estilo_fila(ws3, i, [
            nombre_tipo, datos["total"], round(precio_prom, 2), round(datos["suma_ahorro"], 2)
        ], par=(i % 2 == 0))
    for row in ws3.iter_rows(min_row=4, max_row=len(tipos)+3, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = f'"{moneda}"#,##0.00'
    ajustar_columnas(ws3)
    ws3.freeze_panes = "A4"

    conn.close()

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"PromoIQ_ReporteCompleto_{pais_label}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )



# ── Endpoint: Exportar artículos de UNA tienda, agrupado por ubicación ──
@router.get("/tienda")
def exportar_articulos_tienda(tienda_codigo: str, usuario=Depends(verificar_token)):
    """
    Excel pensado para caminar la tienda físicamente: ordenado por
    ubicación dinámica, para colocar zona por zona sin ir y venir.

    La descripción usa ISNULL/NULLIF como respaldo: si dinamica_detalle.descripcion
    quedó vacía, se completa con la del catálogo maestro por SKU.
    """
    if usuario["rol"] != "tiendas":
        raise HTTPException(status_code=403, detail="Este reporte es exclusivo para el rol Tiendas")

    pais = usuario["pais"]

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT nombre FROM cat_tiendas WHERE codigo = ? AND pais = ? AND activo = 1", (tienda_codigo, pais))
    fila_tienda = cursor.fetchone()
    if not fila_tienda:
        conn.close()
        raise HTTPException(status_code=400, detail="Tienda no válida para tu país")
    tienda_nombre = fila_tienda[0]

    cursor.execute("""
        SELECT dd.sku, ISNULL(NULLIF(dd.descripcion, ''), am.descripcion_cat) as descripcion,
               dd.ubicacion_dinamica, dd.ubicacion_tienda, d.nombre
        FROM dinamica_detalle dd
        INNER JOIN dinamicas d ON dd.dinamica_id = d.id
        LEFT JOIN (
            SELECT sku, pais, MIN(descripcion) as descripcion_cat
            FROM articulos_maestro
            GROUP BY sku, pais
        ) am ON am.sku = dd.sku AND am.pais = dd.pais
        WHERE dd.pais = ? AND dd.tienda = ? AND dd.activo = 1 AND d.activa = 1
        ORDER BY dd.ubicacion_dinamica ASC, dd.sku ASC
    """, (pais, tienda_codigo))
    filas = cursor.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mis Artículos"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    agregar_titulo(
        ws,
        "PromoIQ — Artículos por Ubicación",
        f"Tienda: {tienda_nombre} | Generado el {fecha_reporte}"
    )

    columnas = ["SKU", "Descripción", "Tienda", "Ubicación Dinámica", "Ubicación Tienda", "Campaña"]
    estilo_header(ws, 3, columnas, color_fondo="06847A")
    ws.row_dimensions[3].height = 22

    for i, f in enumerate(filas, 4):
        valores = [
            str(f[0]),
            str(f[1]) if f[1] else "-",
            tienda_nombre,
            str(f[2]) if f[2] else "-",
            str(f[3]) if f[3] else "-",
            str(f[4]),
        ]
        estilo_fila(ws, i, valores, par=(i % 2 == 0))

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"PromoIQ_{tienda_nombre.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )




# ── Endpoint: Exportar comparador de campañas ─────────────
@router.get("/comparador-campanas")
def exportar_comparador_campanas(
    dinamica_id_1: int,
    dinamica_id_2: int,
    pais: str = None,
    usuario=Depends(verificar_token)
):
    """
    Excel con la comparación completa de dos campañas finalizadas —
    resumen de cada una + su top de productos, en hojas separadas.
    """
    if usuario["rol"] not in ["admin", "comercial", "planning", "superadmin"]:
        raise HTTPException(status_code=403, detail="Sin acceso al comparador de campañas")

    if usuario["rol"] == "superadmin":
        if not pais:
            raise HTTPException(status_code=400, detail="Falta indicar el país a consultar")
        pais_ctx = pais
    else:
        pais_ctx = usuario["pais"]

    _verificar_dinamica_finalizada(dinamica_id_1, pais_ctx)
    _verificar_dinamica_finalizada(dinamica_id_2, pais_ctx)

    campana_1 = _calcular_resumen_dinamica(dinamica_id_1, pais_ctx)
    campana_2 = _calcular_resumen_dinamica(dinamica_id_2, pais_ctx)

    conn_moneda = get_connection()
    moneda = obtener_moneda_pais(conn_moneda.cursor(), pais_ctx)
    conn_moneda.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen Comparativo"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    agregar_titulo(
        ws,
        "PromoIQ — Comparador de Campañas",
        f"{campana_1['nombre']} vs. {campana_2['nombre']} | Generado el {fecha_reporte}"
    )

    columnas_resumen = ["Métrica", campana_1['nombre'], campana_2['nombre']]
    estilo_header(ws, 3, columnas_resumen, color_fondo="06847A")
    ws.row_dimensions[3].height = 22

    filas_resumen = [
        ["Vigencia", f"{campana_1['fecha_inicio']} al {campana_1['fecha_fin']}", f"{campana_2['fecha_inicio']} al {campana_2['fecha_fin']}"],
        ["Monto vendido", f"{moneda}{campana_1['resumen']['monto_vendido']:,.2f}", f"{moneda}{campana_2['resumen']['monto_vendido']:,.2f}"],
        ["Unidades vendidas", campana_1['resumen']['unidades_vendidas'], campana_2['resumen']['unidades_vendidas']],
        ["Margen promedio", f"{campana_1['resumen']['margen_promedio']}%", f"{campana_2['resumen']['margen_promedio']}%"],
        ["Cantidad de artículos", campana_1['resumen']['cantidad_articulos'], campana_2['resumen']['cantidad_articulos']],
        ["Cantidad de tiendas", campana_1['resumen']['cantidad_tiendas'], campana_2['resumen']['cantidad_tiendas']],
    ]
    for i, fila in enumerate(filas_resumen, 4):
        estilo_fila(ws, i, fila, par=(i % 2 == 0))

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    # ── Hoja 2: Top productos de la Campaña A ──
    ws2 = wb.create_sheet(title=campana_1['nombre'][:31])
    agregar_titulo(ws2, f"Top productos — {campana_1['nombre']}", f"Generado el {fecha_reporte}")
    columnas_top = ["SKU", "Descripción", "Unidades", "Venta Total", "Margen %"]
    estilo_header(ws2, 3, columnas_top, color_fondo="06847A")
    ws2.row_dimensions[3].height = 22
    for i, p in enumerate(campana_1['top_ventas'], 4):
        estilo_fila(ws2, i, [p['sku'], p['descripcion'], p['unidades'], f"{moneda}{p['venta_total']:,.2f}", f"{p['margen_pct']}%"], par=(i % 2 == 0))
    ajustar_columnas(ws2)
    ws2.freeze_panes = "A4"

    # ── Hoja 3: Top productos de la Campaña B ──
    ws3 = wb.create_sheet(title=campana_2['nombre'][:31])
    agregar_titulo(ws3, f"Top productos — {campana_2['nombre']}", f"Generado el {fecha_reporte}")
    estilo_header(ws3, 3, columnas_top, color_fondo="06847A")
    ws3.row_dimensions[3].height = 22
    for i, p in enumerate(campana_2['top_ventas'], 4):
        estilo_fila(ws3, i, [p['sku'], p['descripcion'], p['unidades'], f"{moneda}{p['venta_total']:,.2f}", f"{p['margen_pct']}%"], par=(i % 2 == 0))
    ajustar_columnas(ws3)
    ws3.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"Comparador_{campana_1['nombre']}_vs_{campana_2['nombre']}.xlsx".replace(" ", "_")

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )


@router.get("/oportunidades-sin-explotar")
def exportar_oportunidades(departamento: str = None, usuario=Depends(verificar_token)):
    """Excel con los códigos nunca promocionados y con venta real, por departamento."""
    if usuario["rol"] != "comercial":
        raise HTTPException(status_code=403, detail="Este reporte es exclusivo para el rol Comercial")

    datos = _calcular_oportunidades_sin_explotar(usuario["pais"], departamento_filtro=departamento)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Oportunidades"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    subtitulo = f"Generado el {fecha_reporte}" + (f" | Filtrado por: {departamento}" if departamento else "")
    agregar_titulo(ws, "PromoIQ — Oportunidades sin Explotar", subtitulo)

    columnas = ["Departamento", "SKU", "Descripción", "Unidades vendidas sin promo"]
    estilo_header(ws, 3, columnas, color_fondo="06847A")
    ws.row_dimensions[3].height = 22

    fila_actual = 4
    for depto in datos["departamentos"]:
        for c in depto["candidatos"]:
            estilo_fila(ws, fila_actual, [depto["departamento"], c["sku"], c["descripcion"], c["unidades"]], par=(fila_actual % 2 == 0))
            fila_actual += 1

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = "Oportunidades_Sin_Explotar" + (f"_{departamento}" if departamento else "") + ".xlsx"
    nombre_archivo = nombre_archivo.replace(" ", "_")

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )



@router.get("/sobre-stock-departamento")
def exportar_sobre_stock_departamento(
    tienda_codigo: str,
    departamento: str = None,
    estado: str = None,
    usuario=Depends(verificar_token)
):
    """
    Excel con el sobre-stock detectado por SKU + tienda, agrupado
    por departamento, con filtros opcionales de estado del artículo.
    """
    if usuario["rol"] != "planning":
        raise HTTPException(status_code=403, detail="Este reporte es exclusivo para el rol Planning")

    datos = _calcular_sobre_stock_departamento(usuario["pais"], tienda_codigo, departamento, estado)

    conn_moneda = get_connection()
    moneda = obtener_moneda_pais(conn_moneda.cursor(), usuario["pais"])
    conn_moneda.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sobre-stock"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    partes_subtitulo = [f"Generado el {fecha_reporte}"]
    if departamento:
        partes_subtitulo.append(f"Departamento: {departamento}")
    if estado:
        partes_subtitulo.append(f"Estado: {estado}")
    subtitulo = " | ".join(partes_subtitulo)

    agregar_titulo(ws, "PromoIQ — Sobre-stock por Tienda", subtitulo)

    columnas = ["Departamento", "SKU", "Descripción", "Estado", "Tienda", "Existencia Actual", "Mínimo", "Máximo", "Exceso en Unidades", "Monto Inmovilizado"]
    estilo_header(ws, 3, columnas, color_fondo="06847A")
    ws.row_dimensions[3].height = 22

    fila_actual = 4
    for depto in datos["departamentos"]:
        for p in depto["productos"]:
            estilo_fila(ws, fila_actual, [
                depto["departamento"], p["sku"], p["descripcion"], p["estado"], p["tienda"],
                p["existencia_actual"], p["minimo"], p["maximo"], p["exceso_unidades"],
                f"{moneda}{p['monto_inmovilizado']:,.2f}"
            ], par=(fila_actual % 2 == 0))
            fila_actual += 1

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    sufijo = ""
    if departamento:
        sufijo += f"_{departamento}"
    if estado:
        sufijo += f"_{estado}"
    nombre_archivo = f"Sobre_Stock{sufijo}_{datetime.now().strftime('%Y%m%d')}.xlsx".replace(" ", "_")

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )


@router.get("/cobertura-tienda")
def exportar_cobertura_tienda(
    tienda_codigo: str,
    departamento: str,
    estados: str,
    usuario=Depends(verificar_token)
):
    """Excel con la cobertura de inventario de un departamento en una tienda, por banda."""
    if usuario["rol"] != "planning":
        raise HTTPException(status_code=403, detail="Este reporte es exclusivo para el rol Planning")

    lista_estados = [e.strip() for e in estados.split(",") if e.strip()]
    if not lista_estados:
        raise HTTPException(status_code=400, detail="Debés seleccionar al menos un estado")

    datos = _calcular_cobertura_tienda(usuario["pais"], tienda_codigo, departamento, lista_estados)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cobertura"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    subtitulo = f"Tienda: {tienda_codigo} | Departamento: {departamento} | Estados: {', '.join(lista_estados)} | Generado el {fecha_reporte}"

    agregar_titulo(ws, "PromoIQ — Cobertura por Departamento", subtitulo)

    columnas = ["SKU", "Descripción", "Estado", "Existencia Actual", "Consumo Mensual", "Cobertura (meses)", "Banda"]
    estilo_header(ws, 3, columnas, color_fondo="06847A")
    ws.row_dimensions[3].height = 22

    for i, p in enumerate(datos["productos"], 4):
        valores = [
            p["sku"], p["descripcion"], p["estado"],
            p["existencia_actual"], p["consumo_mensual"],
            p["cobertura_meses"] if p["cobertura_meses"] is not None else "-",
            p["banda_label"],
        ]
        estilo_fila(ws, i, valores, par=(i % 2 == 0))

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    sufijo_estados = "_" + "_".join(lista_estados)
    nombre_archivo = f"Cobertura_{tienda_codigo}_{departamento}{sufijo_estados}_{datetime.now().strftime('%Y%m%d')}.xlsx".replace(" ", "_")

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )


@router.get("/redistribucion-departamento")
def exportar_redistribucion_departamento(
    departamento: str,
    estados: str,
    usuario=Depends(verificar_token)
):
    """Excel con las sugerencias de redistribución de stock entre tiendas de un departamento."""
    if usuario["rol"] != "planning":
        raise HTTPException(status_code=403, detail="Este reporte es exclusivo para el rol Planning")

    lista_estados = [e.strip() for e in estados.split(",") if e.strip()]
    if not lista_estados:
        raise HTTPException(status_code=400, detail="Debés seleccionar al menos un estado")

    datos = _calcular_redistribucion_departamento(usuario["pais"], departamento, lista_estados)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Redistribución"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    subtitulo = f"Departamento: {departamento} | Estados: {', '.join(lista_estados)} | Generado el {fecha_reporte}"

    agregar_titulo(ws, "PromoIQ — Sugerencias de Redistribución", subtitulo)

    columnas = ["SKU", "Tienda Origen", "Tienda Destino", "Unidades Sugeridas", "Disponible Origen", "Disponible Destino"]
    estilo_header(ws, 3, columnas, color_fondo="B45309")
    ws.row_dimensions[3].height = 22

    for i, s in enumerate(datos["sugerencias"], 4):
        valores = [
            s["sku"], s["tienda_origen"], s["tienda_destino"],
            s["unidades_sugeridas"], s["disponible_origen"], s["disponible_destino"],
        ]
        estilo_fila(ws, i, valores, par=(i % 2 == 0))

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    sufijo_estados = "_" + "_".join(lista_estados)
    nombre_archivo = f"Redistribucion_{departamento}{sufijo_estados}_{datetime.now().strftime('%Y%m%d')}.xlsx".replace(" ", "_")

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )


@router.get("/forecast-demanda")
def exportar_forecast_demanda(
    departamento: str = None,
    estados: str = None,
    codigo_proveedor: str = None,
    usuario=Depends(verificar_token)
):
    """Excel con el forecast de demanda por SKU."""
    if usuario["rol"] != "planning":
        raise HTTPException(status_code=403, detail="Este reporte es exclusivo para el rol Planning")

    lista_estados = [e.strip() for e in estados.split(",") if e.strip()] if estados else None
    datos = _calcular_forecast_demanda(usuario["pais"], departamento, lista_estados, codigo_proveedor)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Forecast"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    partes = [f"Generado el {fecha_reporte}"]
    if departamento:
        partes.append(f"Departamento: {departamento}")
    if estados:
        partes.append(f"Estados: {estados}")
    if codigo_proveedor:
        partes.append(f"Proveedor: {codigo_proveedor}")
    subtitulo = " | ".join(partes)

    agregar_titulo(ws, "PromoIQ — Forecast de Demanda", subtitulo)

    columnas = ["SKU", "Descripción", "Departamento", "Proveedor", "Estado", "Pronóstico Próx. Mes", "Existencia Total", "Cobertura Proyectada (meses)", "Dato Parcial"]
    estilo_header(ws, 3, columnas, color_fondo="1E2761")
    ws.row_dimensions[3].height = 22

    for i, p in enumerate(datos["productos"], 4):
        valores = [
            p["sku"], p["descripcion"], p["departamento"], p["proveedor"], p["estado"],
            p["pronostico_proximo_mes"], p["existencia_total"],
            p["cobertura_proyectada_meses"] if p["cobertura_proyectada_meses"] is not None else "-",
            "Sí" if p["dato_parcial"] else "No",
        ]
        estilo_fila(ws, i, valores, par=(i % 2 == 0))

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"Forecast_Demanda_{datetime.now().strftime('%Y%m%d')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )


@router.get("/impacto-campana-cobertura")
def exportar_impacto_campana_cobertura(
    dinamica_id: int,
    banda: str = None,
    departamento: str = None,
    estados: str = None,
    tienda_codigo: str = None,
    usuario=Depends(verificar_token)
):
    """Excel con la cobertura resultante post-campaña, por SKU+tienda."""
    if usuario["rol"] != "planning":
        raise HTTPException(status_code=403, detail="Este reporte es exclusivo para el rol Planning")

    lista_estados = [e.strip() for e in estados.split(",") if e.strip()] if estados else None
    datos = _calcular_impacto_campana_cobertura(usuario["pais"], dinamica_id, banda, departamento, lista_estados, tienda_codigo)

    conn_moneda = get_connection()
    moneda = obtener_moneda_pais(conn_moneda.cursor(), usuario["pais"])
    conn_moneda.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Impacto Campaña"

    fecha_reporte = datetime.now().strftime("%d/%m/%Y %H:%M")
    partes = [f"{datos['dinamica']['nombre']} ({datos['dinamica']['fecha_inicio']} — {datos['dinamica']['fecha_fin']})"]
    if departamento:
        partes.append(f"Departamento: {departamento}")
    if estados:
        partes.append(f"Estados: {estados}")
    if tienda_codigo:
        partes.append(f"Tienda: {tienda_codigo}")
    if banda:
        partes.append(f"Banda: {banda}")
    partes.append(f"Generado el {fecha_reporte}")
    subtitulo = " | ".join(partes)

    agregar_titulo(ws, "PromoIQ — Impacto de Campaña en Cobertura", subtitulo)

    columnas = ["SKU", "Descripción", "Tienda", "Ubicación Campaña", "Precio Dinámica", "Estado", "Existencia Actual", "Consumo Mensual", "Cobertura (meses)", "Banda"]
    estilo_header(ws, 3, columnas, color_fondo="06847A")
    ws.row_dimensions[3].height = 22

    for i, p in enumerate(datos["productos"], 4):
        valores = [
            p["sku"], p["descripcion"], p["tienda"], p["ubicacion_campana"],
            p["precio_dinamica"], p["estado"], p["existencia_actual"], p["consumo_mensual"],
            p["cobertura_meses"] if p["cobertura_meses"] is not None else "-",
            p["banda_label"],
        ]
        estilo_fila(ws, i, valores, par=(i % 2 == 0))

    for row in ws.iter_rows(min_row=4, max_row=len(datos["productos"]) + 3, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = f'"{moneda}"#,##0.00'

    ajustar_columnas(ws)
    ws.freeze_panes = "A4"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"Impacto_Campana_{datos['dinamica']['nombre']}.xlsx".replace(" ", "_")

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"}
    )
